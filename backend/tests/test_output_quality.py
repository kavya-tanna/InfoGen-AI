"""Regression tests for editorial structure, slide counts and provider fallback."""
import io
import json
from types import SimpleNamespace

import pytest
from pptx import Presentation

from app.config import settings
from app.services.generator import generate_single_format, validate_generated_response
from app.services.normalizer import ContentIntelligence
from app.services.exporters import export_presentation_pptx
from app.services.ai_provider import GeminiProvider


def test_linkedin_has_one_publishable_copy_and_sequential_carousel():
    data = {
        "hook": "120 devices are planned.",
        "post": "120 devices are planned.\n\nA seven-day loan is proposed.\n\nCheck availability.\n\n#Libraries",
        "cta": "Check availability.",
        "hashtags": ["#Libraries", "#libraries"],
        "carousel_slides": [{"slide_number": n, "headline": f"Topic {n}", "body": "Source evidence."} for n in [1, 4, 9]],
    }
    result = validate_generated_response("LinkedIn Post", json.dumps(data), "120 devices are planned. A seven-day loan is proposed. Topic 1. Topic 4. Topic 9.")
    assert result.post.count("120 devices are planned.") == 1
    assert result.post.count("Check availability.") == 1
    assert result.post.casefold().count("#libraries") == 1
    assert [s.slide_number for s in result.carousel_slides] == [1, 2, 3]


@pytest.mark.parametrize("count", [1, 3, 9, 16])
def test_ppt_preserves_ai_slide_count_and_editable_text(count):
    raw = {"title": "Source briefing", "slides": [
        {"slide_number": i*3, "title": f"Specific finding {i}", "key_points": [f"Evidence {i}"],
         "speaker_notes": f"Notes {i}", "layout": "cover" if i == 0 else "editorial"}
        for i in range(count)
    ]}
    source = " ".join(f"Specific finding {i}: Evidence {i}." for i in range(count))
    data = validate_generated_response("Presentation Deck", json.dumps(raw), source)

    assert [s.slide_number for s in data.slides] == list(range(1, count+1))
    prs = Presentation(export_presentation_pptx(data.model_dump()))
    assert len(prs.slides) == count
    for i, slide in enumerate(prs.slides):
        assert any(s.has_text_frame and f"Specific finding {i}" in s.text for s in slide.shapes)
        assert f"Notes {i}" in slide.notes_slide.notes_text_frame.text
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width
            assert shape.top + shape.height <= prs.slide_height


@pytest.mark.parametrize("failure", ["timeout", "malformed", "empty"])
def test_gemini_fallback_after_primary_failure(monkeypatch, failure):
    calls = []

    class Primary:
        name = "nvidia"
        model_name = "primary"

        def generate(self, prompt, **kwargs):
            calls.append(("primary", kwargs["timeout_seconds"]))
            if failure == "timeout":
                raise TimeoutError()
            return "not JSON" if failure == "malformed" else "{}"

    class Fallback:
        name = "gemini"
        model_name = "fallback"

        def generate(self, prompt, **kwargs):
            calls.append(("fallback", kwargs["timeout_seconds"]))
            return '{"headline":"Library pilot","key_takeaways":["120 devices are planned."]}'

    monkeypatch.setattr("app.services.generator.get_provider", lambda: Primary())
    monkeypatch.setattr("app.services.generator.get_fallback_provider", lambda name: Fallback())
    result = generate_single_format("Executive Summary", ContentIntelligence(source_text="120 devices are planned."))
    assert result.generation_mode == "ai" and result.provider == "gemini"
    assert result.model == "fallback"
    assert any("Gemini fallback" in warning for warning in result.warnings)
    assert calls[-1][0] == "fallback"
    assert all(0 < timeout <= settings.ai_timeout_seconds for _, timeout in calls)


def test_gemini_adapter_sends_json_and_timeout():
    captured = {}

    def create(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(choices=[SimpleNamespace(finish_reason="stop", message=SimpleNamespace(content='{"ok":true}'))])

    provider = object.__new__(GeminiProvider)
    provider._model = "gemini-3.6-flash"
    provider._client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    assert json.loads(provider.generate("Return JSON", timeout_seconds=21)) == {"ok": True}
    assert captured["timeout"] == 21
    assert captured["response_format"] == {"type": "json_object"}


def test_x_removes_duplicate_primary_and_number_prefixes():
    result = validate_generated_response("Twitter / X Post", json.dumps({
        "primary_post": "A new pilot is planned.",
        "thread": ["1/ A new pilot is planned.", "2/ Devices are lent for seven days."],
    }), "A new pilot is planned. Devices are lent for seven days.")
    assert result.thread == ["Devices are lent for seven days."]


@pytest.mark.parametrize("name,data", [
    ("Presentation Deck", {"title": "Title", "slides": [{"title": "Empty"}]}),
    ("Infographic", {"title": "Title", "sections": [{"section_title": "Empty"}]}),
    ("Executive Summary", []),
])
def test_empty_structure_is_not_accepted_as_ai(name, data):
    with pytest.raises(ValueError):
        validate_generated_response(name, json.dumps(data), "A source document.")


def test_new_calculated_number_is_rejected():
    with pytest.raises(ValueError, match="unsupported numerical"):
        validate_generated_response("Presentation Deck", json.dumps({
            "title": "Library pilot", "slides": [{
                "title": "Pilot scope", "key_points": ["120 devices across four branches."],
                "takeaway": "20 devices per branch weekly.",
            }],
        }), "120 devices across four branches over six weeks.")


def test_spelled_source_numbers_allow_numeric_output():
    result = validate_generated_response("LinkedIn Post", '{"post":"4 branches, 7-day loans."}', "Four branches, seven-day loans.")
    assert result.post == "4 branches, 7-day loans."


def test_synthetic_label_is_kept_in_publishable_copy():
    result = validate_generated_response("LinkedIn Post", '{"post":"A library pilot is planned."}', "Synthetic news article for demonstration. A fictional library pilot is planned.")
    assert result.post.startswith("Synthetic example:")


def test_long_social_paragraph_gets_readable_breaks():
    sentences = ["The library will provide orientation and erase borrower data after each return."] * 7
    result = validate_generated_response("LinkedIn Post", json.dumps({"post": " ".join(sentences)}), " ".join(sentences))
    assert "\n\n" in result.post
    assert result.post.replace("\n\n", " ") == " ".join(sentences)
