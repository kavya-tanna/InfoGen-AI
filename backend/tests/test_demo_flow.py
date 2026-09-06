import io
import json
import re
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.schemas.outputs import OUTPUT_FORMAT_MAP, OUTPUT_KEY_MAP

client = TestClient(app)
samples = json.loads((Path(__file__).resolve().parents[2] / "samples/demo.json").read_text())


@pytest.mark.parametrize("sample", samples, ids=lambda s: s["id"])
def test_sample_all_formats_and_exports(sample):
    response = client.post("/api/generate", json={
        "source": sample["content"], "content_type": sample["content_type"],
        "output_format": ", ".join(OUTPUT_FORMAT_MAP),
    })
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["metadata"]["persisted"]
    assert result["metadata"]["content_type"] == sample["content_type"]
    outputs = result["outputs"]
    assert set(outputs) == set(OUTPUT_KEY_MAP.values())
    summary = outputs["executive_summary"]
    assert summary["headline"] == sample["title"]
    assert all(p in sample["content"] for p in summary["key_takeaways"])
    assert all(o["generation_mode"] == "extractive" for o in outputs.values())
    assert "100% Grounded" not in result["output"]
    assert "Readiness score 94%" not in result["output"]
    if sample["id"] == "report":
        assert "72%" in [s["value"] for s in outputs["infographic"]["key_statistics"]]
    assert len(outputs["twitter_post"]["primary_post"]) <= 280
    assert re.search(r"00:\d{2}:\d{2}\.000 --> 00:\d{2}:\d{2}\.000", outputs["video_package"]["vtt"])
    assert client.get(f"/api/jobs/{result['job_id']}").json()["status"] == "completed"
    assert client.get(f"/api/outputs/{result['job_id']}").json()["outputs"] == outputs
    for key, signature in [("executive_summary", b"%PDF"), ("advisory_document", b"%PDF"),
                           ("presentation_deck", b"PK"), ("infographic", b"<!DOCTYPE")]:
        exported = client.post("/api/export", json={"format_type": key, "data": outputs[key]})
        assert exported.status_code == 200, exported.text[:300]
        assert exported.content.startswith(signature)
        if key == "presentation_deck":
            from pptx import Presentation
            assert len(Presentation(io.BytesIO(exported.content)).slides) >= 5


@pytest.mark.parametrize("extension", ["txt", "pdf", "docx"])
def test_upload(extension):
    text = samples[0]["content"]
    if extension == "txt":
        data = text.encode()
    elif extension == "pdf":
        import fitz
        doc = fitz.open()
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(30, 30, 565, 800), text, fontsize=10)
        data = doc.tobytes()
        doc.close()
    else:
        from docx import Document
        doc = Document()
        doc.add_paragraph(text)
        buffer = io.BytesIO()
        doc.save(buffer)
        data = buffer.getvalue()
    response = client.post("/api/generate", files={"files": (f"sample.{extension}", data)},
                           data={"output_format": "Executive Summary", "content_type": "News Article"})
    assert response.status_code == 200, response.text
    assert "Riverton" in response.json()["outputs"]["executive_summary"]["headline"]


@pytest.mark.parametrize("body", [[], {"source": 12}, {"source": "x" * 50001}, {"source": "abc", "tone": []}])
def test_bad_json_input(body):
    assert client.post("/api/generate", json=body).status_code == 400


@pytest.mark.parametrize("name,data", [("bad.pdf", b"not a pdf"), ("empty.txt", b""),
                                      ("image.png", b"image"), ("binary.txt", b"a\x00b")])
def test_bad_upload(name, data):
    response = client.post("/api/generate", files={"files": (name, data)})
    assert response.status_code == 400
    assert response.json()["error"]["message"]


@pytest.mark.parametrize("failure", ["timeout", "bad-json", "empty-json", "invented-ioc"])
def test_provider_failure_recovers_to_source(monkeypatch, failure):
    class BrokenProvider:
        name = "test-provider"

        def generate(self, *args, **kwargs):
            if failure == "timeout":
                raise TimeoutError()
            if failure == "bad-json":
                return "not json"
            if failure == "empty-json":
                return "{}"
            return '{"headline":"CVE-2099-99999","key_takeaways":["Invented identifier"]}'
    monkeypatch.setattr("app.services.generator.get_provider", lambda: BrokenProvider())
    response = client.post("/api/generate", json={"source": samples[0]["content"]})
    assert response.status_code == 200
    output = response.json()["outputs"]["executive_summary"]
    assert output["generation_mode"] == "extractive"
    assert "AI generation failed" in output["warnings"][0]
    assert "CVE-2099" not in json.dumps(output)


def test_missing_risk_is_not_invented():
    output = client.post("/api/generate", json={"source": "A library opened a reading room on Monday."}).json()["outputs"]["executive_summary"]
    assert output["risks"] == ["Not available in source."]
    assert output["priority"] == "UNSPECIFIED"


def test_export_escapes_user_markup():
    result = client.post("/api/export", json={"format_type": "infographic",
        "data": {"title": '<script>alert("test")</script>', "sections": []}})
    assert result.status_code == 200
    assert "<script>" not in result.text
    assert "&lt;script&gt;" in result.text
    pdf = client.post("/api/export", json={"format_type": "executive_summary",
        "data": {"headline": "Research: A < B & C", "context": "<img src='file:///unknown'>"}})
    assert pdf.status_code == 200


def test_unknown_job_and_invalid_url():
    assert client.get("/api/jobs/unknown").status_code == 404
    result = client.post("/api/generate", json={"source": "http://127.0.0.1/private"})
    assert result.status_code == 400
    assert result.json()["error"]["code"] == "URL_EXTRACTION_FAILED"
